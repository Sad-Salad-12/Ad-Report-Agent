"""Deterministic transformations from classified source tables to report JSON."""

from __future__ import annotations

import math
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Sequence, Tuple

from openpyxl.utils import get_column_letter

from .history import HistoryStore
from .ingest import IngestionError, header_position, ingest_bundle, row_value
from .models import (
    AudienceMetrics,
    CampaignMetrics,
    CanonicalReport,
    CreativeMetrics,
    DailyMetrics,
    DataPipelineResult,
    HistoryRecord,
    IngestedBundle,
    KeywordMetrics,
    LineageEntry,
    OverallMetrics,
    ProductAnalysis,
    ProductConfig,
    ProductSummary,
    ReportConfig,
    SourceCellRef,
    SourcePeriod,
    SourceRow,
    SourceTable,
    TrafficMetrics,
    load_report_config,
)


class TransformationError(ValueError):
    """Raised when canonical data cannot be produced without an assumption."""


AVERAGE_PURCHASE_VALUE_HEADERS = (
    "Website purchase average order value (USD)",
    "BRAND A: website Purchase average order value (USD)",
    # Backward compatibility with historical exports whose vendor prefix is embedded
    # in the source column name.
    "Rokid: website Purchase 客单价 (USD)",
)


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _number(value: Any, context: str) -> float:
    if isinstance(value, bool) or _blank(value):
        raise TransformationError(f"{context}: expected a number, got {value!r}")
    if isinstance(value, (int, float)):
        result = float(value)
    elif isinstance(value, str):
        text = value.strip().replace(",", "").replace("%", "")
        negative = text.startswith("(") and text.endswith(")")
        if negative:
            text = text[1:-1]
        text = re.sub(r"^[\$€£]", "", text).strip()
        try:
            result = float(text)
        except ValueError as exc:
            raise TransformationError(f"{context}: expected a number, got {value!r}") from exc
        if negative:
            result = -result
    else:
        raise TransformationError(f"{context}: expected a number, got {type(value).__name__}")
    if not math.isfinite(result):
        raise TransformationError(f"{context}: number must be finite")
    return result


def _integer(value: Any, context: str) -> int:
    result = _number(value, context)
    rounded = round(result)
    if not math.isclose(result, rounded, abs_tol=1e-9):
        raise TransformationError(f"{context}: expected an integer, got {result}")
    return int(rounded)


def _optional_integer(value: Any, context: str) -> Optional[int]:
    return None if _blank(value) else _integer(value, context)


def _date(value: Any, context: str) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        match = re.search(r"20\d{2}-\d{2}-\d{2}", value)
        if match:
            return date.fromisoformat(match.group(0))
    raise TransformationError(f"{context}: expected an ISO date, got {value!r}")


def _normalized_text(value: Any) -> str:
    return re.sub(
        r"\s+",
        " ",
        unicodedata.normalize("NFKC", str(value or "")),
    ).strip().casefold()


def _match_product(text: Any, products: Sequence[ProductConfig], context: str) -> ProductConfig:
    normalized = _normalized_text(text)
    matches: List[Tuple[int, ProductConfig]] = []
    for product in products:
        for alias in product.aliases:
            normalized_alias = _normalized_text(alias)
            if normalized_alias and normalized_alias in normalized:
                matches.append((len(normalized_alias), product))
    if not matches:
        raise TransformationError(f"{context}: no configured product alias matched {text!r}")
    longest = max(length for length, _ in matches)
    winners = {product.id: product for length, product in matches if length == longest}
    if len(winners) != 1:
        raise TransformationError(f"{context}: product alias match is ambiguous for {text!r}")
    return next(iter(winners.values()))


def _cell_ref(table: SourceTable, row: SourceRow, header: str) -> SourceCellRef:
    try:
        column, raw_header = header_position(table, header)
        raw_value = row_value(row, header)
    except IngestionError as exc:
        raise TransformationError(str(exc)) from exc
    return SourceCellRef(
        source_kind="bundle",
        file=table.file_name,
        sheet=table.sheet_name,
        row=row.row_number,
        column=column,
        cell=f"{get_column_letter(column)}{row.row_number}",
        header=raw_header,
        raw_value=raw_value,
    )


def _first_present_header(table: SourceTable, candidates: Sequence[str]) -> str:
    """Return the first supported alias actually present in a source table."""

    for candidate in candidates:
        try:
            _, raw_header = header_position(table, candidate)
            return raw_header
        except IngestionError:
            continue
    joined = ", ".join(candidates)
    raise TransformationError(
        f"{table.file_name}: required header not found; accepted aliases: {joined}"
    )


def _add_source_lineage(
    lineage: List[LineageEntry],
    target: str,
    table: SourceTable,
    row: SourceRow,
    header: str,
    transform: str = "identity",
) -> None:
    lineage.append(
        LineageEntry(
            target=target,
            transform=transform,
            sources=[_cell_ref(table, row, header)],
        )
    )


def _add_formula_lineage(
    lineage: List[LineageEntry],
    target: str,
    table: SourceTable,
    row: SourceRow,
    headers: Sequence[str],
    transform: str,
) -> None:
    lineage.append(
        LineageEntry(
            target=target,
            transform=transform,
            sources=[_cell_ref(table, row, header) for header in headers],
        )
    )


def _one(rows: Sequence[SourceRow], context: str) -> SourceRow:
    if len(rows) != 1:
        raise TransformationError(f"{context}: expected exactly one row, found {len(rows)}")
    return rows[0]


def _overall(
    table: SourceTable,
    config: ReportConfig,
    lineage: List[LineageEntry],
) -> OverallMetrics:
    row = _one(table.rows, "overall")
    if table.period_start is None or table.period_end is None:
        raise TransformationError("overall: missing Reporting starts/ends")
    headers = {
        "amount_spent": "Amount spent (USD)",
        "purchase_value": "Purchases conversion value",
        "purchase_roas": "Purchase ROAS (return on ad spend)",
        "purchases": "Purchases",
        "adds_to_cart": "Adds to cart",
        "cost_per_purchase": "Cost per purchase",
        "cost_per_add_to_cart": "Cost per add to cart",
        "link_clicks": "Link clicks",
        "ctr": "CTR (link click-through rate)",
    }
    values = {key: row_value(row, header) for key, header in headers.items()}
    adds = _integer(values["adds_to_cart"], "overall.adds_to_cart")
    clicks = _integer(values["link_clicks"], "overall.link_clicks")
    if clicks <= 0:
        raise TransformationError("overall.link_clicks must be positive")
    metrics = OverallMetrics(
        market=config.market,
        period_start=table.period_start,
        period_end=table.period_end,
        amount_spent=_number(values["amount_spent"], "overall.amount_spent"),
        purchase_value=_number(values["purchase_value"], "overall.purchase_value"),
        purchase_roas=_number(values["purchase_roas"], "overall.purchase_roas"),
        purchases=_integer(values["purchases"], "overall.purchases"),
        adds_to_cart=adds,
        cost_per_purchase=_number(values["cost_per_purchase"], "overall.cost_per_purchase"),
        cost_per_add_to_cart=_number(
            values["cost_per_add_to_cart"], "overall.cost_per_add_to_cart"
        ),
        link_clicks=clicks,
        ctr=_number(values["ctr"], "overall.ctr"),
        add_to_cart_rate=adds / clicks * 100.0,
    )
    for field, header in headers.items():
        _add_source_lineage(lineage, f"current_overall.{field}", table, row, header)
    _add_formula_lineage(
        lineage,
        "current_overall.add_to_cart_rate",
        table,
        row,
        ["Adds to cart", "Link clicks"],
        "adds_to_cart / link_clicks * 100 (percentage points)",
    )
    _add_source_lineage(
        lineage, "current_overall.period_start", table, row, "Reporting starts", "ISO date"
    )
    _add_source_lineage(
        lineage, "current_overall.period_end", table, row, "Reporting ends", "ISO date"
    )
    return metrics


def _daily(table: SourceTable, lineage: List[LineageEntry]) -> List[DailyMetrics]:
    result: List[Tuple[DailyMetrics, SourceRow]] = []
    for row in table.rows:
        item = DailyMetrics(
            day=_date(row_value(row, "Day"), f"by_day row {row.row_number}.day"),
            adds_to_cart=_integer(row_value(row, "Adds to cart"), "daily.adds_to_cart"),
            purchases=_integer(row_value(row, "Purchases"), "daily.purchases"),
            purchase_roas=_number(
                row_value(row, "Purchase ROAS (return on ad spend)"), "daily.purchase_roas"
            ),
            amount_spent=_number(row_value(row, "Amount spent (USD)"), "daily.amount_spent"),
            purchase_value=_number(
                row_value(row, "Purchases conversion value"), "daily.purchase_value"
            ),
        )
        result.append((item, row))
    result.sort(key=lambda pair: pair[0].day)
    headers = {
        "day": "Day",
        "adds_to_cart": "Adds to cart",
        "purchases": "Purchases",
        "purchase_roas": "Purchase ROAS (return on ad spend)",
        "amount_spent": "Amount spent (USD)",
        "purchase_value": "Purchases conversion value",
    }
    for index, (_, row) in enumerate(result):
        for field, header in headers.items():
            _add_source_lineage(lineage, f"daily.{index}.{field}", table, row, header)
    return [item for item, _ in result]


def _product_rows(
    table: SourceTable,
    config: ReportConfig,
    text_header: str,
    *,
    include: Optional[Any] = None,
) -> Dict[str, SourceRow]:
    mapped: Dict[str, SourceRow] = {}
    for row in table.rows:
        if include is not None and not include(row):
            continue
        text = row_value(row, text_header, required=False)
        if _blank(text):
            continue
        product = _match_product(text, config.products, f"{table.kind} row {row.row_number}")
        if product.id in mapped:
            raise TransformationError(
                f"{table.kind}: more than one row matched product {product.id}"
            )
        mapped[product.id] = row
    missing = [product.id for product in config.products if product.id not in mapped]
    if missing:
        raise TransformationError(f"{table.kind}: missing configured products {missing}")
    return mapped


def _product_summary(
    product: ProductConfig,
    table: SourceTable,
    row: SourceRow,
    index: int,
    lineage: List[LineageEntry],
) -> ProductSummary:
    item = ProductSummary(
        product_id=product.id,
        display_name=product.display_name,
        campaign_name=str(row_value(row, "Campaign name")).strip(),
        adds_to_cart=_integer(row_value(row, "Adds to cart"), "product.adds_to_cart"),
        purchases=_integer(row_value(row, "Purchases"), "product.purchases"),
        purchase_roas=_number(
            row_value(row, "Purchase ROAS (return on ad spend)"), "product.purchase_roas"
        ),
        average_purchase_value=_number(
            row_value(row, "Average purchases conversion value"),
            "product.average_purchase_value",
        ),
        purchase_value=_number(
            row_value(row, "Purchases conversion value"), "product.purchase_value"
        ),
        amount_spent=_number(row_value(row, "Amount spent (USD)"), "product.amount_spent"),
    )
    headers = {
        "campaign_name": "Campaign name",
        "adds_to_cart": "Adds to cart",
        "purchases": "Purchases",
        "purchase_roas": "Purchase ROAS (return on ad spend)",
        "average_purchase_value": "Average purchases conversion value",
        "purchase_value": "Purchases conversion value",
        "amount_spent": "Amount spent (USD)",
    }
    for field, header in headers.items():
        _add_source_lineage(lineage, f"product_summaries.{index}.{field}", table, row, header)
        _add_source_lineage(
            lineage, f"product_analyses.{index}.summary.{field}", table, row, header
        )
    return item


def _campaign(
    product: ProductConfig,
    table: SourceTable,
    row: SourceRow,
    index: int,
    lineage: List[LineageEntry],
) -> CampaignMetrics:
    average_purchase_value_header = _first_present_header(
        table, AVERAGE_PURCHASE_VALUE_HEADERS
    )
    item = CampaignMetrics(
        product_id=product.id,
        campaign_name=str(row_value(row, "Campaign name")).strip(),
        amount_spent=_number(row_value(row, "Amount spent (USD)"), "campaign.amount_spent"),
        purchase_value=_number(
            row_value(row, "Purchases conversion value"), "campaign.purchase_value"
        ),
        purchase_roas=_number(
            row_value(row, "Purchase ROAS (return on ad spend)"), "campaign.purchase_roas"
        ),
        purchases=_integer(row_value(row, "Purchases"), "campaign.purchases"),
        adds_to_cart=_integer(row_value(row, "Adds to cart"), "campaign.adds_to_cart"),
        cost_per_purchase=_number(
            row_value(row, "Cost per purchase"), "campaign.cost_per_purchase"
        ),
        cost_per_add_to_cart=_number(
            row_value(row, "Cost per add to cart"), "campaign.cost_per_add_to_cart"
        ),
        purchase_rate=_number(
            row_value(row, "Purchases rate per landing page views"), "campaign.purchase_rate"
        )
        * 100.0,
        add_to_cart_rate=_number(
            row_value(row, "Add to cart Rate"), "campaign.add_to_cart_rate"
        )
        * 100.0,
        impressions=_integer(row_value(row, "Impressions"), "campaign.impressions"),
        link_clicks=_integer(row_value(row, "Link clicks"), "campaign.link_clicks"),
        ctr=_number(row_value(row, "CTR (link click-through rate)"), "campaign.ctr"),
        cpm=_number(
            row_value(row, "CPM (cost per 1,000 impressions)"), "campaign.cpm"
        ),
        cpc=_number(row_value(row, "CPC (cost per link click)"), "campaign.cpc"),
        average_purchase_value=_number(
            row_value(row, average_purchase_value_header),
            "campaign.average_purchase_value",
        ),
    )
    headers = {
        "campaign_name": "Campaign name",
        "amount_spent": "Amount spent (USD)",
        "purchase_value": "Purchases conversion value",
        "purchase_roas": "Purchase ROAS (return on ad spend)",
        "purchases": "Purchases",
        "adds_to_cart": "Adds to cart",
        "cost_per_purchase": "Cost per purchase",
        "cost_per_add_to_cart": "Cost per add to cart",
        "purchase_rate": "Purchases rate per landing page views",
        "add_to_cart_rate": "Add to cart Rate",
        "impressions": "Impressions",
        "link_clicks": "Link clicks",
        "ctr": "CTR (link click-through rate)",
        "cpm": "CPM (cost per 1,000 impressions)",
        "cpc": "CPC (cost per link click)",
        "average_purchase_value": average_purchase_value_header,
    }
    for field, header in headers.items():
        transform = "multiply by 100 (ratio to percentage points)" if field in {
            "purchase_rate",
            "add_to_cart_rate",
        } else "identity"
        _add_source_lineage(
            lineage,
            f"product_analyses.{index}.campaign.{field}",
            table,
            row,
            header,
            transform,
        )
    return item


def _creative(
    product: ProductConfig,
    table: SourceTable,
    row: SourceRow,
    index: int,
    lineage: List[LineageEntry],
) -> CreativeMetrics:
    average_purchase_value_header = _first_present_header(
        table, AVERAGE_PURCHASE_VALUE_HEADERS
    )
    item = CreativeMetrics(
        product_id=product.id,
        ad_name=str(row_value(row, "Ads")).strip(),
        campaign_name=str(row_value(row, "Campaign name")).strip(),
        amount_spent=_number(row_value(row, "Amount spent (USD)"), "creative.amount_spent"),
        cpm=_number(row_value(row, "CPM (cost per 1,000 impressions)"), "creative.cpm"),
        cpc=_number(row_value(row, "CPC (cost per link click)"), "creative.cpc"),
        ctr=_number(row_value(row, "CTR (link click-through rate)"), "creative.ctr"),
        cost_per_add_to_cart=_number(
            row_value(row, "Cost per add to cart"), "creative.cost_per_add_to_cart"
        ),
        cost_per_purchase=_number(
            row_value(row, "Cost per purchase"), "creative.cost_per_purchase"
        ),
        purchase_roas=_number(
            row_value(row, "Purchase ROAS (return on ad spend)"), "creative.purchase_roas"
        ),
        average_purchase_value=_number(
            row_value(row, average_purchase_value_header),
            "creative.average_purchase_value",
        ),
        adds_to_cart=_integer(row_value(row, "Adds to cart"), "creative.adds_to_cart"),
        purchases=_integer(row_value(row, "Purchases"), "creative.purchases"),
    )
    headers = {
        "ad_name": "Ads",
        "campaign_name": "Campaign name",
        "amount_spent": "Amount spent (USD)",
        "cpm": "CPM (cost per 1,000 impressions)",
        "cpc": "CPC (cost per link click)",
        "ctr": "CTR (link click-through rate)",
        "cost_per_add_to_cart": "Cost per add to cart",
        "cost_per_purchase": "Cost per purchase",
        "purchase_roas": "Purchase ROAS (return on ad spend)",
        "average_purchase_value": average_purchase_value_header,
        "adds_to_cart": "Adds to cart",
        "purchases": "Purchases",
    }
    for field, header in headers.items():
        _add_source_lineage(
            lineage, f"product_analyses.{index}.creative.{field}", table, row, header
        )
    return item


def _traffic(
    table: SourceTable, config: ReportConfig, lineage: List[LineageEntry]
) -> List[TrafficMetrics]:
    result: List[TrafficMetrics] = []
    seen: set[str] = set()
    for index, row in enumerate(table.rows):
        campaign_name = str(row_value(row, "Campaign name")).strip()
        product = _match_product(campaign_name, config.products, f"traffic row {row.row_number}")
        if product.id in seen:
            raise TransformationError(f"traffic: duplicate product row {product.id}")
        seen.add(product.id)
        result.append(
            TrafficMetrics(
                product_id=product.id,
                campaign_name=campaign_name,
                amount_spent=_number(row_value(row, "Amount spent (USD)"), "traffic.amount_spent"),
                purchases=_optional_integer(row_value(row, "Purchases"), "traffic.purchases"),
                adds_to_cart=_optional_integer(
                    row_value(row, "Adds to cart"), "traffic.adds_to_cart"
                ),
                impressions=_integer(row_value(row, "Impressions"), "traffic.impressions"),
                link_clicks=_integer(row_value(row, "Link clicks"), "traffic.link_clicks"),
                ctr=_number(row_value(row, "CTR (link click-through rate)"), "traffic.ctr"),
                cpm=_number(
                    row_value(row, "CPM (cost per 1,000 impressions)"), "traffic.cpm"
                ),
                cpc=_number(row_value(row, "CPC (cost per link click)"), "traffic.cpc"),
                landing_page_views=_integer(
                    row_value(row, "Landing page views"), "traffic.landing_page_views"
                ),
                cost_per_landing_page_view=_number(
                    row_value(row, "Cost per landing page view"),
                    "traffic.cost_per_landing_page_view",
                ),
                landing_page_view_rate=_number(
                    row_value(row, "Landing page views rate per link clicks"),
                    "traffic.landing_page_view_rate",
                ),
            )
        )
        headers = {
            "campaign_name": "Campaign name",
            "amount_spent": "Amount spent (USD)",
            "purchases": "Purchases",
            "adds_to_cart": "Adds to cart",
            "impressions": "Impressions",
            "link_clicks": "Link clicks",
            "ctr": "CTR (link click-through rate)",
            "cpm": "CPM (cost per 1,000 impressions)",
            "cpc": "CPC (cost per link click)",
            "landing_page_views": "Landing page views",
            "cost_per_landing_page_view": "Cost per landing page view",
            "landing_page_view_rate": "Landing page views rate per link clicks",
        }
        for field, header in headers.items():
            _add_source_lineage(lineage, f"traffic.{index}.{field}", table, row, header)
    missing = [product.id for product in config.products if product.id not in seen]
    if missing:
        raise TransformationError(f"traffic: missing configured products {missing}")
    return result


def _audience(
    table: SourceTable, config: ReportConfig, lineage: List[LineageEntry]
) -> List[AudienceMetrics]:
    result: List[AudienceMetrics] = []
    seen: set[str] = set()
    for index, row in enumerate(table.rows):
        ad_set_name = str(row_value(row, "Ad Set Name")).strip()
        product = _match_product(ad_set_name, config.products, f"audience row {row.row_number}")
        if product.id in seen:
            raise TransformationError(f"audience: duplicate product row {product.id}")
        seen.add(product.id)
        result.append(
            AudienceMetrics(
                product_id=product.id,
                ad_set_name=ad_set_name,
                adds_to_cart=_integer(row_value(row, "Adds to cart"), "audience.adds_to_cart"),
                purchases=_integer(row_value(row, "Purchases"), "audience.purchases"),
                purchase_roas=_number(
                    row_value(row, "Purchase ROAS (return on ad spend)"),
                    "audience.purchase_roas",
                ),
                purchase_value=_number(
                    row_value(row, "Purchases conversion value"), "audience.purchase_value"
                ),
                amount_spent=_number(
                    row_value(row, "Amount spent (USD)"), "audience.amount_spent"
                ),
            )
        )
        headers = {
            "ad_set_name": "Ad Set Name",
            "adds_to_cart": "Adds to cart",
            "purchases": "Purchases",
            "purchase_roas": "Purchase ROAS (return on ad spend)",
            "purchase_value": "Purchases conversion value",
            "amount_spent": "Amount spent (USD)",
        }
        for field, header in headers.items():
            _add_source_lineage(lineage, f"audience.{index}.{field}", table, row, header)
    missing = [product.id for product in config.products if product.id not in seen]
    if missing:
        raise TransformationError(f"audience: missing configured products {missing}")
    return result


def _clean_keyword(value: Any) -> str:
    text = str(value or "").strip()
    while len(text) >= 2 and text[0] == text[-1] == '"':
        text = text[1:-1].strip()
    if not text:
        raise TransformationError("keyword cannot be blank")
    return text


def _keywords(
    table: SourceTable, top_n: int, lineage: List[LineageEntry]
) -> List[KeywordMetrics]:
    ranked: List[Tuple[float, int, SourceRow]] = []
    for row in table.rows:
        conversions = _number(row_value(row, "Conversions"), "keyword.conversions")
        ranked.append((conversions, row.row_number, row))
    ranked.sort(key=lambda item: (-item[0], item[1]))
    selected = ranked[:top_n]
    if len(selected) < top_n:
        raise TransformationError(
            f"keyword: expected at least {top_n} data rows, found {len(selected)}"
        )
    result: List[KeywordMetrics] = []
    headers = {
        "keyword": "Keyword",
        "conversions": "Conversions",
        "conversion_value_per_cost": "Conv. value / cost",
        "cost": "Cost",
        "conversion_value": "Conv. value",
    }
    for index, (conversions, _, row) in enumerate(selected):
        result.append(
            KeywordMetrics(
                rank=index + 1,
                keyword=_clean_keyword(row_value(row, "Keyword")),
                conversions=conversions,
                conversion_value_per_cost=_number(
                    row_value(row, "Conv. value / cost"),
                    "keyword.conversion_value_per_cost",
                ),
                cost=_number(row_value(row, "Cost"), "keyword.cost"),
                conversion_value=_number(
                    row_value(row, "Conv. value"), "keyword.conversion_value"
                ),
            )
        )
        for field, header in headers.items():
            transform = "strip enclosing quote characters" if field == "keyword" else "identity"
            _add_source_lineage(
                lineage, f"keywords.{index}.{field}", table, row, header, transform
            )
    return result


def _history_lineage(lineage: List[LineageEntry], record: HistoryRecord) -> None:
    for field, value in record.overall.model_dump(mode="json").items():
        lineage.append(
            LineageEntry(
                target=f"previous_overall.{field}",
                transform="latest history record with period_end before current period_start",
                sources=[
                    SourceCellRef(
                        source_kind=record.source_kind,
                        file=record.source_file,
                        header=field,
                        raw_value=value,
                    )
                ],
            )
        )


def transform_bundle(
    bundle: IngestedBundle,
    config: ReportConfig,
    previous_overall: Optional[OverallMetrics] = None,
    *,
    previous_history: Optional[HistoryRecord] = None,
) -> CanonicalReport:
    """Build the complete data model used by all eight sample slides."""

    lineage: List[LineageEntry] = []
    overall_table = bundle.tables["overall"]
    current_overall = _overall(overall_table, config, lineage)
    daily = _daily(bundle.tables["by_day"], lineage)

    by_product = bundle.tables["by_product"]
    summary_rows = _product_rows(by_product, config, "Name")
    product_summaries = [
        _product_summary(product, by_product, summary_rows[product.id], index, lineage)
        for index, product in enumerate(config.products)
    ]

    campaign_table = bundle.tables["campaign"]
    campaign_rows = _product_rows(
        campaign_table,
        config,
        "Campaign name",
        include=lambda row: not _blank(
            row_value(row, "Purchase ROAS (return on ad spend)", required=False)
        ),
    )

    creative_table = bundle.tables["creative"]
    pinned_rows: Dict[str, SourceRow] = {}
    for product in config.products:
        matches = [
            row
            for row in creative_table.rows
            if _normalized_text(row_value(row, "Ads")) == _normalized_text(product.creative_pin)
        ]
        pinned_rows[product.id] = _one(matches, f"creative pin {product.creative_pin!r}")
        creative_product = _match_product(
            row_value(pinned_rows[product.id], "Campaign name"),
            config.products,
            f"creative pin {product.creative_pin!r}",
        )
        if creative_product.id != product.id:
            raise TransformationError(
                f"creative pin {product.creative_pin!r} belongs to {creative_product.id}, not {product.id}"
            )

    product_analyses: List[ProductAnalysis] = []
    for index, product in enumerate(config.products):
        product_analyses.append(
            ProductAnalysis(
                product_id=product.id,
                display_name=product.display_name,
                creative_asset=product.creative_asset,
                summary=product_summaries[index],
                campaign=_campaign(
                    product, campaign_table, campaign_rows[product.id], index, lineage
                ),
                creative=_creative(
                    product, creative_table, pinned_rows[product.id], index, lineage
                ),
            )
        )

    traffic = _traffic(bundle.tables["traffic_campaign"], config, lineage)
    audience = _audience(bundle.tables["audience"], config, lineage)
    keywords = _keywords(bundle.tables["keyword"], config.keyword_top_n, lineage)

    history_record = previous_history
    if history_record is not None:
        previous_overall = history_record.overall
        _history_lineage(lineage, history_record)

    return CanonicalReport(
        schema_version=config.schema_version,
        template_version=config.template_version,
        market=config.market,
        country_label=config.country_label,
        currency=config.currency,
        report_title=config.report_title,
        period_start=current_overall.period_start,
        period_end=current_overall.period_end,
        source_bundle_sha256=bundle.bundle_sha256,
        source_files={kind: table.file_name for kind, table in bundle.tables.items()},
        source_periods=[
            SourcePeriod(
                source_kind=kind,
                file_name=table.file_name,
                period_start=table.period_start,
                period_end=table.period_end,
            )
            for kind, table in bundle.tables.items()
        ],
        current_overall=current_overall,
        previous_overall=previous_overall,
        daily=daily,
        product_summaries=product_summaries,
        product_analyses=product_analyses,
        traffic=traffic,
        audience=audience,
        keywords=keywords,
        lineage=lineage,
    )


def _resolve_config_path(config_path: Path, configured: Optional[str]) -> Optional[Path]:
    if not configured:
        return None
    path = Path(configured)
    if path.is_absolute():
        return path
    base = config_path.parent.parent if config_path.parent.name == "config" else config_path.parent
    return base / path


def run_data_pipeline(
    bundle_path: str | Path,
    config_path: str | Path,
    *,
    history_db_path: str | Path | None = None,
    seed_path: str | Path | None = None,
    ingested_bundle: IngestedBundle | None = None,
) -> DataPipelineResult:
    """Ingest, transform, attach prior history, and validate in one stable call."""

    from .validation import validate_report

    config_file = Path(config_path).resolve()
    config = load_report_config(config_file)
    resolved_db = (
        Path(history_db_path)
        if history_db_path is not None
        else _resolve_config_path(config_file, config.history.database)
    )
    if resolved_db is None:
        raise TransformationError("history database path is not configured")
    resolved_seed = (
        Path(seed_path)
        if seed_path is not None
        else _resolve_config_path(config_file, config.history.seed_file)
    )
    bundle = ingested_bundle if ingested_bundle is not None else ingest_bundle(bundle_path)
    overall_table = bundle.tables["overall"]
    if overall_table.period_start is None:
        raise TransformationError("overall source does not declare a reporting period")
    history = HistoryStore(resolved_db, resolved_seed)
    previous_record = history.get_previous_record(config.market, overall_table.period_start)
    report = transform_bundle(bundle, config, previous_history=previous_record)
    validation = validate_report(report, config, report.previous_overall)
    return DataPipelineResult(
        report=report,
        validation=validation,
        previous_overall=report.previous_overall,
    )
